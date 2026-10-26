"""
Extracts the per-item criteria of the multi-criteria triangulation ranking from the
working spreadsheet into a plain CSV that is versioned with the paper.

Sources (read only), all in Gdrive_files/Item Reduction/:
  1. "Ranking revisited considering reliabitiy- loadings- and IRT.xlsx", tab "3 fators merged"
     (April 2026, one row per item, 49 items): values, ranks, IRT scores, and the first
     triangulated priority -> column `priority` (used to build Model_3 of Versions A-D).
  2. "Ranking revisited 2 - AfterHandpickedItemsbyAFandP- considering reliabitiy- loadings- and IRT.xlsx",
     same tab: the priority revised by Petra Isenberg and Anne-Flore Cabouat when they chose the
     candidate pool of the combinatorial search (handpicked_items_list_by_factor.csv; history in
     "Historic of handpicked_items_list_by_factor.pdf") -> column `priority_final`.
  3. "ranking_table_full.xlsx": the considerations completed by Eliane on 2026-10-20 with the reasons
     given by Petra and Anne-Flore -> column `considerations` (when the cell is filled).

Output:
  phase2-2-item-reduction/04_ranking/ranking_criteria.csv

Run once from the repository root:
  python phase2-2-item-reduction/04_ranking/extract_ranking_criteria.py
The CSV is the input of build_ranking_table.py; the spreadsheet is not needed afterwards.
"""
from pathlib import Path
import csv
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
GD = ROOT / "Gdrive_files" / "Item Reduction"
SRC = GD / "Ranking revisited considering reliabitiy- loadings- and IRT.xlsx"
SRC_FINAL = GD / "Ranking revisited 2 - AfterHandpickedItemsbyAFandP- considering reliabitiy- loadings- and IRT.xlsx"
SRC_NOTES = GD / "ranking_table_full.xlsx"
OUT = Path(__file__).resolve().parent / "ranking_criteria.csv"

FACTORS = {"System": "system", "Visualization": "visualization", "Task_cognition": "task_cognition"}

# column positions in tab "3 fators merged" (0-based), checked against the header row
COLS = {
    "factor": (0, "Factor"),
    "rank_reliability": (1, "Reliability  ranking"),
    "alpha_if_dropped": (3, "raw_alpha"),
    "rank_loading": (4, "Factor loadings rank"),
    "loading": (6, "Factor loadings"),
    "rank_irt": (7, "IRT ranking"),
    "priority": (9, "Ranking revisited considering reliabitiy - loadings and IRT"),
    "considerations": (11, "Considerations "),
    "code": (13, "Item ID"),
    "irt_monotonicity": (14, "Wave Succession"),
    "irt_discrimination": (15, "Inclination of Curve/Peaks Height "),
    "irt_middle_category": (16, "Neutral Point Centering"),
    "irt_coverage": (17, "Full Scale Utilization - Range (-3 to +3)"),
}


def read_tab(path, header_required=True):
    """Rows of tab "3 fators merged". The header row is found and checked; in the revised
    spreadsheet it was overwritten, so there the columns are taken by position and every value
    is compared with the April spreadsheet in main()."""
    ws = openpyxl.load_workbook(path, data_only=True)["3 fators merged"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows[:10]) if r[0] == "Factor"), None)
    if hdr is None:
        if header_required:
            sys.exit(f"Header row not found in {path.name}")
        print(f"NOTE {path.name[:30]}...: no header row; columns read by position")
    else:
        for key, (pos, name) in COLS.items():
            if rows[hdr][pos] != name:
                sys.exit(f"{path.name}: unexpected header in column {pos}: {rows[hdr][pos]!r} (expected {name!r})")
    start = 0 if hdr is None else hdr + 1
    return [r for r in rows[start:] if r[0] in FACTORS]


def as_int(v):
    return "" if v in (None, "no classified") else int(v)


def main():
    rows = []
    for r in read_tab(SRC):
        rec = {k: r[pos] for k, (pos, _) in COLS.items()}
        rec["factor"] = FACTORS[rec["factor"]]
        for k in ("rank_reliability", "rank_loading", "rank_irt", "priority",
                  "irt_monotonicity", "irt_discrimination", "irt_middle_category", "irt_coverage"):
            rec[k] = as_int(rec[k])
        rec["considerations"] = (rec["considerations"] or "").strip()
        rows.append(rec)
    if len(rows) != 49:
        sys.exit(f"Expected 49 items, found {len(rows)}")
    by_code = {r["code"]: r for r in rows}

    # revised (final) priority; every other column must be unchanged
    pos_pri, pos_code = COLS["priority"][0], COLS["code"][0]
    final = {r[pos_code]: r for r in read_tab(SRC_FINAL, header_required=False)}
    if set(final) != set(by_code):
        sys.exit(f"{SRC_FINAL.name}: items differ: {sorted(set(final) ^ set(by_code))}")
    for code, r in final.items():
        by_code[code]["priority_final"] = as_int(r[pos_pri])
        for key, (pos, _) in COLS.items():
            if key in ("priority", "factor"):
                continue
            a, b = by_code[code][key], r[pos]
            b = as_int(b) if isinstance(a, int) or a == "" and key != "considerations" else b
            if key == "considerations":
                b = (b or "").strip()
            if a != b and not (isinstance(a, float) and abs(a - b) < 1e-9):
                print(f"NOTE {code} {key}: {a!r} (April) vs {b!r} ({SRC_FINAL.name[:20]}...)")

    # completed considerations
    wb = openpyxl.load_workbook(SRC_NOTES, data_only=True).active
    head = [c.value for c in wb[1]]
    ci, cc, cp = head.index("code"), head.index("considerations"), head.index("priority")
    for r in wb.iter_rows(min_row=2, values_only=True):
        if r[ci] in by_code:
            if r[cc]:
                by_code[r[ci]]["considerations"] = str(r[cc]).strip()
            p = r[cp]
            p = "" if p is None or isinstance(p, str) else int(p)
            if p != by_code[r[ci]]["priority_final"]:
                print(f"NOTE {r[ci]}: priority {r[cp]!r} in {SRC_NOTES.name} vs "
                      f"{by_code[r[ci]]['priority_final']!r} in {SRC_FINAL.name[:20]}... (the latter is used)")

    fields = [k for k in COLS if k != "considerations"] + ["priority_final", "considerations"]
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {len(rows)} items to {OUT}")


if __name__ == "__main__":
    main()
